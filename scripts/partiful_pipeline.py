"""Partiful event -> LinkedIn-enriched workbook, in one command.

    python partiful_pipeline.py --event https://partiful.com/e/<ID> \
        --out "~/Desktop/<Event> Guest List - LinkedIn.xlsx"

Does the whole pipeline with no agent round-trips:

  1. refresh the stored Firebase token
  2. getEventInfo + getGuests (hard-fails on masked names)
  3. filter to real first+last names (drops single names, initial-surnames,
     couples, company names)
  4. Exa searches in parallel (default 8 workers), with a
     `site:linkedin.com/in` retry for names that surface no profile
  5. score candidates per profile slug and build the workbook
  6. audit: assert every URL cell is backed by a literal /in/ hit
  7. optionally create a standalone Google Sheet (--upload)
  8. optionally add a tab to the Oxen-Forecast master sheet (--append-master)

Auth comes from <hermes>/partiful-auth.json; the Exa key from <hermes>/.env.
Exit code is non-zero on any failure, so a cron job can alert on it.
"""
from __future__ import annotations

import argparse
import difflib
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
import unicodedata
from datetime import datetime
from collections import Counter
from concurrent.futures import ThreadPoolExecutor

# openpyxl is imported lazily inside write_workbook()/upload_to_sheets() so that
# partiful_incremental.py can `from partiful_pipeline import classify` without
# needing openpyxl installed in the interpreter doing the scrape.

HERMES = os.path.join(os.environ.get("LOCALAPPDATA", os.path.expanduser("~")), "hermes")
AUTH = os.path.join(HERMES, "partiful-auth.json")
ENV = os.path.join(HERMES, ".env")
API = "https://api.partiful.com/"
MASTER_ID = "1XNtycJwuDnl_OuPw_1KAASU19TRD2WstEiYH9JPFy7o"
MAX_TAB_NAME_ATTEMPTS = 50
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/155.0.0.0 Safari/537.36")


class PipelineError(Exception):
    """Raised instead of sys.exit so embedders (GUI, cron) can catch failures."""

    def __init__(self, msg, code=1):
        super().__init__(msg)
        self.code = code


def die(msg, code=1):
    raise PipelineError(msg, code)


# Embedding hooks: a GUI or scheduler sets LOG_FN (and optionally PROGRESS)
# to receive every status line / enrichment tick in addition to stdout.
LOG_FN = None
PROGRESS = None


def log(*args):
    msg = " ".join(str(a) for a in args)
    print(msg)
    if LOG_FN:
        LOG_FN(msg)


# ---------------------------------------------------------------- Partiful

def post(url, data, headers):
    req = urllib.request.Request(url, data=data, headers=headers, method="POST")
    with urllib.request.urlopen(req, timeout=45) as r:
        return r.read().decode()


def refresh_access(cfg):
    body = urllib.parse.urlencode({"grant_type": "refresh_token",
                                   "refresh_token": cfg["refresh_token"]}).encode()
    try:
        out = json.loads(post(
            f"https://securetoken.googleapis.com/v1/token?key={cfg['firebase_api_key']}",
            body, {"Content-Type": "application/x-www-form-urlencoded",
                   "Referer": "https://partiful.com/"}))
    except Exception as e:
        die(f"token refresh failed ({e}). The refresh token was likely revoked - "
            "log into Partiful in Chrome and re-extract it.")
    cfg["access_token"] = out["id_token"]
    cfg["user_id"] = out["user_id"]
    if out.get("refresh_token"):
        cfg["refresh_token"] = out["refresh_token"]
    json.dump(cfg, open(AUTH, "w", encoding="utf-8"))
    return cfg["access_token"], cfg["user_id"]


def api(fn, params, access, uid):
    payload = json.dumps({"data": {"params": params, "userId": uid}}).encode()
    body = post(API + fn, payload, {
        "Content-Type": "application/json", "Origin": "https://partiful.com",
        "Referer": "https://partiful.com/", "Authorization": "Bearer " + access,
        "User-Agent": UA})
    return json.loads(body)["result"]["data"]


MASKED = re.compile(r"^[A-Za-z]\s*x{2,}")
TOKCH = re.compile(r"^[A-Za-z\u00c0-\u024f'\u2019.\-]+$")
INITIAL = re.compile(r"^[A-Za-z]?\.?$")
COMPANY = re.compile(r"\b(inc|llc|ltd|corp|company|group|labs?|studio|ventures|capital"
                     r"|partners|naturals|team)\b", re.I)


def strip_handles(n):
    n = re.sub(r"@\S+", "", n)
    n = re.sub(r"https?://\S+", "", n)
    n = re.sub(r"[\U0001F000-\U0001FAFF\u2600-\u27BF\uFE0F]", "", n)
    return re.sub(r"\s+", " ", n).strip(" ,;-|")


def classify(guests):
    passing, dropped, seen = [], [], set()
    for g in guests:
        raw = (g.get("user") or {}).get("name") or g.get("name")
        if not raw:
            continue
        key = raw.strip().lower()
        if key in seen:
            continue
        seen.add(key)
        n = strip_handles(raw)
        if not n:
            dropped.append([raw, "no name after cleanup"]); continue
        if COMPANY.search(n):
            dropped.append([raw, "company/brand-like"]); continue
        if re.search(r"\s*(?:&|\band\b)\s*", n, re.I):
            dropped.append([raw, "multiple people (couple/duo)"]); continue
        toks = n.split()
        # only a *surname* initial disqualifies; a middle initial is fine
        real = [t for t in toks if not INITIAL.match(t)]
        if len(real) < 2:
            dropped.append([raw, "initial instead of a surname" if len(toks) > 1
                            else "first name only"]); continue
        if INITIAL.match(toks[-1]):
            dropped.append([raw, "initial instead of a surname"]); continue
        if len(toks) >= 2 and all(TOKCH.match(t) for t in toks):
            passing.append({"name": n, "status": g.get("status"),
                            "plusOneCount": g.get("plusOneCount", 0),
                            "rsvpDate": g.get("rsvpDate")})
        elif len(toks) == 1:
            dropped.append([raw, "first name only"])
        else:
            dropped.append([raw, "not a two-part name"])
    return passing, dropped


# --------------------------------------------------------------------- Exa

def load_exa_key():
    if not os.path.isfile(ENV):
        die(f"no .env at {ENV}")
    for line in open(ENV, encoding="utf-8", errors="ignore"):
        if line.startswith("EXA_API_KEY="):
            return line.split("=", 1)[1].strip().strip('"').strip("'")
    die("EXA_API_KEY not found in .env")


def exa_search(key, query, n=5, tries=3):
    payload = json.dumps({"query": query, "numResults": n,
                          "contents": {"text": {"maxCharacters": 1500}}}).encode()
    last = None
    for attempt in range(tries):
        try:
            req = urllib.request.Request(
                "https://api.exa.ai/search", data=payload, method="POST",
                headers={"Content-Type": "application/json", "x-api-key": key})
            with urllib.request.urlopen(req, timeout=60) as r:
                return json.loads(r.read().decode()).get("results", [])
        except Exception as e:
            last = e
            time.sleep(1.5 * (attempt + 1))
    return [{"error": str(last)[:200]}]


def norm_hits(results):
    out = []
    for h in results:
        out.append({"url": h.get("url"), "title": h.get("title") or "",
                    "desc": (h.get("text") or h.get("summary") or "")[:1500]})
    return out


def enrich(key, names, workers, city, on_progress=None):
    store, t0 = {}, time.time()

    def one(n):
        hits = norm_hits(exa_search(key, f'"{n}" LinkedIn {city}'))
        if not any(re.search(r"linkedin\.com/in/", h.get("url") or "") for h in hits):
            hits += norm_hits(exa_search(key, f'"{n}" site:linkedin.com/in'))
        return n, hits

    done = 0
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for n, hits in ex.map(one, names):
            store[n] = hits
            done += 1
            if done % 25 == 0:
                log(f"    searched {done}/{len(names)}  ({time.time()-t0:.0f}s)")
            if on_progress:
                on_progress(done, len(names))
            if PROGRESS:
                PROGRESS(done, len(names))
    return store, time.time() - t0


# ---------------------------------------------------------------- scoring

METRO_RE = re.compile(r"([A-Z][A-Za-z.'\- ]+Metropolitan Area)\s*\(US\)")
CITY_RE = re.compile(
    r"([A-Z][A-Za-z.'\-]+(?: [A-Z][A-Za-z.'\-]+)*,\s*"
    r"(?:California|CA|New York|NY|Texas|TX|Washington|WA|Illinois|IL|Massachusetts|MA|"
    r"Florida|FL|Georgia|GA|Colorado|CO|Arizona|AZ|Nevada|NV|Oregon|OR|Utah|UT|Michigan|MI|"
    r"Ohio|OH|Pennsylvania|PA|North Carolina|NC|Virginia|VA|Tennessee|TN|Minnesota|MN|"
    r"Missouri|MO|Maryland|MD|Indiana|IN|Wisconsin|WI|Connecticut|CT|New Jersey|NJ|"
    r"[A-Z][a-z]+(?: [A-Z][a-z]+)?),\s*United States)")

LA_CORE = [
 "los angeles","santa monica","culver city","venice","west hollywood","beverly hills","pasadena",
 "burbank","glendale","long beach","manhattan beach","torrance","el segundo","sherman oaks",
 "monterey park","santa clarita","inglewood","redondo beach","hermosa beach","playa vista",
 "woodland hills","encino","north hollywood","studio city","van nuys","chatsworth","canoga park",
 "reseda","tarzana","calabasas","agoura hills","malibu","pacific palisades","brentwood","westwood",
 "century city","hollywood","echo park","silver lake","los feliz","koreatown","marina del rey",
 "westchester","gardena","hawthorne","lawndale","carson","compton","pomona","claremont","whittier",
 "norwalk","downey","cerritos","lakewood","alhambra","arcadia","monrovia","azusa","covina",
 "el monte","rosemead","san gabriel","temple city","south pasadena","san marino","la canada",
 "la crescenta","montrose","eagle rock","highland park","atwater village","boyle heights",
 "lincoln heights","el sereno","west adams","leimert park","crenshaw","baldwin hills",
 "huntington park","south gate","lynwood","paramount","bell gardens","maywood","glendora",
 "san dimas","la verne","diamond bar","walnut","rowland heights","hacienda heights","west covina",
 "baldwin park","duarte","altadena","sierra madre","san pedro","wilmington","rancho palos verdes",
 "palos verdes","lomita","signal hill","bellflower","artesia","la mirada","pico rivera",
 "santa fe springs","universal city","toluca lake","tujunga","sunland","sylmar","san fernando",
 "pac oima","arleta","panorama city","northridge","granada hills","porter ranch","north hills",
 "winnetka","sherman village","topanga","westlake","greater los angeles","los angeles county",
]
LA_ADJACENT = [
 "thousand oaks","newbury park","westlake village","simi valley","oxnard","ventura","camarillo",
 "moorpark","riverside","san bernardino","ontario, ca","rancho cucamonga","fontana","moreno valley",
 "corona, ca","chino","chino hills","redlands","temecula","murrieta","lancaster","palmdale",
]
NON_LA = [
 "irvine","anaheim","tustin","orange, ca","orange county","santa ana","newport beach","costa mesa",
 "huntington beach","fullerton","garden grove","mission viejo","laguna","yorba linda","brea",
 "buena park","fountain valley","westminster","seal beach","placentia","la habra","san clemente",
 "dana point","aliso viejo","lake forest","rancho santa margarita","san francisco","san jose",
 "sunnyvale","palo alto","mountain view","santa clara","cupertino","fremont","oakland","berkeley",
 "pleasanton","san mateo","redwood city","menlo park","san bruno","milpitas","hayward","union city",
 "los altos","campbell","saratoga","walnut creek","concord, ca","san diego","santa barbara",
 "sacramento","new york","nyc","brooklyn","manhattan","austin","seattle","chicago","boston","miami",
 "atlanta","denver","phoenix","scottsdale","portland","salt lake","detroit","columbus",
 "philadelphia","pittsburgh","raleigh","nashville","minneapolis","kansas city","baltimore",
 "indianapolis","milwaukee","hartford","newark, nj","jersey city","london","toronto","vancouver",
 "dubai","singapore","bangalore","bengaluru","tel aviv","paris","berlin","houston","dallas",
 "san antonio","orlando","tampa","charlotte","st. louis","st louis","cincinnati","cleveland",
 "buffalo","providence","richmond","new orleans","las vegas","honolulu","anchorage","stockholm",
 "sweden","united kingdom","germany","france","canada","india","israel","australia","netherlands",
 "spain","italy","japan","china","brazil","mexico","ireland","switzerland","denmark","norway",
 "finland","poland","portugal","austria","belgium","czech","romania","ukraine","turkey",
]
ORDER = {"Yes": 0, "Adjacent": 1, "Unknown": 2, "No": 3}


def norm(s):
    return re.sub(r"\s+", " ", re.sub(r"[.\u2019']", "", (s or "").lower())).strip()


def name_forms(n):
    f = norm(n)
    forms = {f}
    parts = f.split()
    if len(parts) > 2:
        forms.add(parts[0] + " " + parts[-1])
    return forms


def locs_in(text):
    out = [m.group(1).strip() for m in METRO_RE.finditer(text)]
    out += [m.group(1).strip() for m in CITY_RE.finditer(text)]
    return out


def classify_loc(loc):
    L = (loc or "").lower()
    for k in LA_CORE:
        if k in L:
            return "Yes"
    for k in LA_ADJACENT:
        if k in L:
            return "Adjacent"
    for k in NON_LA:
        if k in L:
            return "No"
    if "california" in L or ", ca" in L:
        return "Unknown"
    return "No"


def slug_from(url):
    """ONLY a literal /in/ URL proves a profile slug - a post URL's slug is the
    author's, and for a company post it is the company. Never construct one."""
    m = re.search(r"linkedin\.com/in/([A-Za-z0-9\-_%.]+)", url or "")
    return (m.group(1), True) if m else (None, False)


def name_tokens(s):
    """Lowercase word tokens of a name or profile title (hyphens kept)."""
    return {t for t in re.split(r"[^A-Za-z\u00c0-\u024f'\-]+", (s or "").lower()) if t}


def title_matches(title, name):
    """True when every meaningful token of `name` appears as a token of the
    profile's OWN title.

    Token boundaries matter, not substrings: a plain `in` test lets `Seth Tam`
    match the title `Seth Tams`, which is a different person one character away.
    """
    want = [t for t in name_tokens(name) if len(t) > 1]
    have = name_tokens(title)
    return all(t in have for t in want)


def ordered_tokens(s):
    """Lowercase word tokens, in order (hyphens kept)."""
    return [t for t in re.split(r"[^A-Za-z\u00c0-\u024f'\-]+", (s or "").lower()) if t]


def _flat(s):
    """Accent-folded, punctuation-free lowercase form (for slug comparison)."""
    s = unicodedata.normalize("NFKD", (s or "").lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", s)


def surname_ok(title, name, slug=None):
    """Does this candidate carry evidence of the person's surname?

    The surname is the discriminator: a title that merely contains the *given*
    name is a different person (`Adem Betta` must not match `ALLEN ADEM`).

    Three kinds of evidence count:
      - the surname as a title token, accent-folded (`radim vysoký` ~ `Vysoky`)
      - a hyphenated surname whole or part by part (`Anisa Anorve-Gomez` ->
        `Anissa Gomez`; `Tosin Ade-Ajayi` keeps its hyphen)
      - the surname inside the profile *slug*, which is name-derived and
        survives an abbreviated title (`Nour H.` for `/in/nourhassoun`)

    Near-variants are only allowed on long tokens, and slug evidence only on
    surnames of 5+ characters. Fuzzy-matching short surnames re-admits exactly
    the near-miss class the skill warns about: `Seth Tam` -> `Seth Tams`
    scores 0.857, so a bare 0.85 ratio keeps a different person.
    """
    nt = [t for t in ordered_tokens(name) if len(t) > 1]
    if not nt:
        return False
    surname = nt[-1]
    probes = [surname] + [p for p in surname.split("-") if len(p) > 1]
    have = ordered_tokens(title)
    for p in probes:
        fp = _flat(p)
        for h in have:
            if p == h or fp == _flat(h):
                return True
            if len(p) >= 6 and len(h) >= 6 and difflib.SequenceMatcher(None, p, h).ratio() >= 0.85:
                return True
    if slug:
        fs = _flat(slug)
        for p in probes:
            if len(_flat(p)) >= 5 and _flat(p) in fs:
                return True
    return False


def strict_identity(rows, store):
    """Blank the URL on rows whose matched profile shows no surname evidence.

    Returns the names that were blanked. Rows that keep their URL but fail the
    full-name check stay flagged for review - those are usually nicknames.
    """
    blanked = []
    for r in rows:
        if not r["url"]:
            continue
        slug = r["url"].rsplit("/", 1)[-1]
        titles = [h.get("title") or "" for h in store.get(r["name"], [])
                  if re.search(r"linkedin\.com/in/" + re.escape(slug), h.get("url") or "")]
        if titles and not any(surname_ok(t, r["name"], slug) for t in titles):
            blanked.append(r["name"])
            r.update(url=None, loc=None, flag="Unknown", conf="Low",
                             note=f"No LinkedIn match - nearest hit was a different person "
                                  f"({titles[0][:40]})", headline="")
    return blanked


def build_rows(store):
    rows = []
    for name in sorted(store, key=str.lower):
        forms = name_forms(name)
        cands = {}
        for h in store[name]:
            slug, is_profile = slug_from(h.get("url"))
            if not slug:
                continue
            blob = (h.get("title") or "") + " \n " + (h.get("desc") or "")
            name_hit = any(f in norm(blob) for f in forms)
            c = cands.setdefault(slug, {"slug": slug, "profile": False, "name_hit": False,
                                        "title_hit": False, "locs": [], "flags": []})
            c["profile"] = c["profile"] or is_profile
            c["name_hit"] = c["name_hit"] or name_hit
            c["title_hit"] = c["title_hit"] or title_matches(h.get("title") or "", name)
            if not c.get("title") and (h.get("title") or ""):
                c["title"] = h["title"]
            for l in locs_in(blob):
                if l not in c["locs"]:
                    c["locs"].append(l)
                    c["flags"].append(classify_loc(l))
        ranked = []
        for c in cands.values():
            if not c["profile"]:
                continue
            flag = sorted(c["flags"], key=lambda f: ORDER[f])[0] if c["flags"] else "Unknown"
            loc = next((l for l, f in zip(c["locs"], c["flags"]) if f == flag), None)
            score = ((5 if c["name_hit"] else 0) + (4 if c["title_hit"] else 0)
                     + 3 + (2 if c["locs"] else 0))
            ranked.append({**c, "flag": flag, "loc": loc, "score": score})
        ranked.sort(key=lambda c: (-c["score"], ORDER[c["flag"]]))
        n = len(ranked)
        if not ranked:
            rows.append({"name": name, "url": None, "loc": None, "flag": "Unknown",
                                 "conf": "Low", "note": "No LinkedIn profile surfaced", "n": 0,
                                 "headline": ""})
            continue
        best = ranked[0]
        if best["flag"] == "Yes":
            conf = "High" if n == 1 else "Medium"
        elif best["flag"] == "Adjacent":
            conf = "Medium" if n == 1 else "Low"
        else:
            conf = "Low"
        note = f"{n} candidate{'s' if n != 1 else ''}"
        if best["loc"]:
            note += f"; {best['loc']}"
        if not best["title_hit"]:
            # The provenance audit passes this row happily - the slug really did
            # come from a literal /in/ hit. Only the identity check notices that
            # the matched profile's own title does not carry this person's name.
            conf = "Low"
            note = "NAME CHECK: matched profile title lacks this name | " + note
        rows.append({"name": name, "url": f"https://linkedin.com/in/{best['slug']}",
                     "loc": best["loc"], "flag": best["flag"], "conf": conf,
                     "note": note, "n": n, "headline": best.get("title") or ""})
    return rows


def audit(rows, store):
    bad = []
    for r in rows:
        if not r["url"]:
            continue
        slug = r["url"].rsplit("/", 1)[-1]
        if not any(re.search(r"linkedin\.com/in/" + re.escape(slug), h.get("url") or "")
                   for h in store.get(r["name"], [])):
            bad.append((r["name"], slug))
    return bad


# -------------------------------------------------------------- workbook

def write_workbook(rows, dropped, ev, eid, out, extra_meta=()):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1F3864")

    ws = wb.active
    ws.title = "LA-Area"
    cols = ["Name", "LinkedIn Profile", "Location", "LA Area?", "Confidence", "Notes"]
    ws.append(cols)
    ordered = sorted(rows, key=lambda r: (ORDER[r["flag"]], r["name"].lower()))
    for r in ordered:
        ws.append([r["name"], r["url"], r["loc"], r["flag"], r["conf"], r["note"]])
    for i, r in enumerate(ordered, start=2):
        if r["url"]:
            c = ws.cell(i, 2)
            c.hyperlink = r["url"]
            c.font = Font(color="0563C1", underline="single")
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c)
        cell.font = hf
        cell.fill = hfill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    for col, w in zip("ABCDEF", (30, 46, 42, 11, 12, 54)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Not Included")
    ws2.append(["Name", "Reason"])
    for name, reason in dropped:
        ws2.append([name, reason])
    for c in (1, 2):
        cell = ws2.cell(1, c)
        cell.font = hf
        cell.fill = hfill
    ws2.freeze_panes = "A2"
    ws2.column_dimensions["A"].width = 46
    ws2.column_dimensions["B"].width = 34

    ws3 = wb.create_sheet("Event")
    conf = Counter(r["conf"] for r in rows)
    meta = [
        ("Event", ev.get("title")),
        ("Partiful URL", f"https://partiful.com/e/{eid}"),
        ("Start (UTC)", ev.get("startDate")),
        ("Venue", ev.get("location")),
        ("Names in this build", len(rows)),
        ("With LinkedIn URL", sum(1 for r in rows if r["url"])),
        ("LA Area? = Yes", sum(1 for r in rows if r["flag"] == "Yes")),
        ("Confidence High", conf["High"]),
        ("Confidence Medium", conf["Medium"]),
        ("Confidence Low", conf["Low"]),
        ("Source", "Partiful getGuests API (authenticated) + Exa search API"),
    ]
    for k, v in list(meta) + list(extra_meta):
        ws3.append([k, v])
    for row in ws3.iter_rows(min_row=1, max_row=ws3.max_row, min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 95

    wb.save(out)


def upload_to_sheets(xlsx_path, title=None):
    """Create a Google Sheet from the workbook, mirroring every tab.

    Imports are lazy so the pipeline still runs on an interpreter without the
    Google libraries when --upload is not requested.
    """
    import openpyxl
    from google.oauth2.credentials import Credentials
    from googleapiclient.discovery import build

    token = os.path.join(HERMES, "google_token.json")
    if not os.path.isfile(token):
        die(f"not authenticated with Google (no {token}). "
            "Run the google-workspace setup first.")
    creds = Credentials.from_authorized_user_file(token)
    sh = build("sheets", "v4", credentials=creds)

    wb = openpyxl.load_workbook(xlsx_path)
    tabs = []
    for ws in wb.worksheets:
        rows = [[("" if c is None else c) for c in r]
                for r in ws.iter_rows(values_only=True)]
        while rows and all(c == "" for c in rows[-1]):
            rows.pop()
        tabs.append({"name": ws.title, "rows": rows})

    created = sh.spreadsheets().create(body={
        "properties": {"title": title or os.path.splitext(os.path.basename(xlsx_path))[0]},
        "sheets": [{"properties": {"title": t["name"]}} for t in tabs],
    }).execute()
    sid = created["spreadsheetId"]
    # sheetId is assigned by the API, not 0..n - read it back from the response
    for t, s in zip(tabs, created["sheets"]):
        t["sheet_id"] = s["properties"]["sheetId"]

    reqs = []
    for t in tabs:
        if t["rows"]:
            sh.spreadsheets().values().update(
                spreadsheetId=sid, range=f"'{t['name']}'!A1",
                valueInputOption="RAW", body={"values": t["rows"]}).execute()
        ncols = max((len(r) for r in t["rows"]), default=1)
        reqs.append({"updateSheetProperties": {
            "properties": {"sheetId": t["sheet_id"],
                           "gridProperties": {"frozenRowCount": 1}},
            "fields": "gridProperties.frozenRowCount"}})
        if len(t["rows"]) > 1:
            reqs.append({"setBasicFilter": {"filter": {
                "range": {"sheetId": t["sheet_id"], "startRowIndex": 0,
                          "endRowIndex": len(t["rows"]),
                          "startColumnIndex": 0, "endColumnIndex": ncols}}}})
        reqs.append({"repeatCell": {
            "range": {"sheetId": t["sheet_id"], "startRowIndex": 0, "endRowIndex": 1,
                      "startColumnIndex": 0, "endColumnIndex": ncols},
            "cell": {"userEnteredFormat": {
                "textFormat": {"bold": True,
                               "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}},
                "backgroundColor": {"red": 0.122, "green": 0.220, "blue": 0.392}}},
            "fields": "userEnteredFormat(textFormat,backgroundColor)"}})
    sh.spreadsheets().batchUpdate(spreadsheetId=sid, body={"requests": reqs}).execute()
    return f"https://docs.google.com/spreadsheets/d/{sid}"


def append_to_master(rows, tab_name, spreadsheet_id=MASTER_ID):
    """Add a new tab to a spreadsheet with enriched rows (per-event tab).

    Defaults to the Oxen-Forecast master sheet. If the tab name already
    exists, retries with " (2)", " (3)", ... suffixes.
    Returns the URL to the new tab.
    """
    from google.oauth2.credentials import Credentials
    from googleapiclient.discovery import build
    from googleapiclient.errors import HttpError

    token = os.path.join(HERMES, "google_token.json")
    if not os.path.isfile(token):
        die(f"not authenticated with Google (no {token}). "
            "Run the google-workspace setup first.")
    creds = Credentials.from_authorized_user_file(token)
    sh = build("sheets", "v4", credentials=creds)

    base = (tab_name or "Event")[:80].strip() or "Event"
    tab_name = base
    attempt_n = 1
    new_sheet_id = None
    while True:
        try:
            resp = sh.spreadsheets().batchUpdate(
                spreadsheetId=spreadsheet_id,
                body={"requests": [{"addSheet": {
                    "properties": {"title": tab_name}}}]}
            ).execute()
            new_sheet_id = resp["replies"][0]["addSheet"]["properties"]["sheetId"]
            break
        except HttpError as e:
            if "already exists" not in str(e):
                raise
            attempt_n += 1
            if attempt_n > MAX_TAB_NAME_ATTEMPTS:
                die(f"Could not find a free tab name for {base!r} in master "
                    f"sheet after {MAX_TAB_NAME_ATTEMPTS} attempts.")
            tab_name = f"{base[:75]} ({attempt_n})"

    SHEET_COLS = ["Name", "LinkedIn Profile", "Location",
                  "LA Area?", "Confidence", "Notes"]
    sheet_rows = [SHEET_COLS] + [
        [r["name"], r["url"] or "", r["loc"] or "",
         r["flag"], r["conf"], r["note"]]
        for r in sorted(rows, key=lambda x: (ORDER[x["flag"]], x["name"].lower()))
    ]

    if sheet_rows:
        sh.spreadsheets().values().update(
        spreadsheetId=spreadsheet_id,
            range=f"'{tab_name}'!A1",
            valueInputOption="RAW",
            body={"values": sheet_rows}).execute()

    ncols = max((len(r) for r in sheet_rows), default=6)
    nrows = len(sheet_rows)
    style_reqs = [
        {"updateSheetProperties": {
            "properties": {"sheetId": new_sheet_id,
                           "gridProperties": {"frozenRowCount": 1}},
            "fields": "gridProperties.frozenRowCount"}},
        {"repeatCell": {
            "range": {"sheetId": new_sheet_id,
                      "startRowIndex": 0, "endRowIndex": 1,
                      "startColumnIndex": 0, "endColumnIndex": ncols},
            "cell": {"userEnteredFormat": {
                "textFormat": {"bold": True,
                               "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}},
                "backgroundColor": {"red": 0.122, "green": 0.220, "blue": 0.392}}},
            "fields": "userEnteredFormat(textFormat,backgroundColor)"}},
    ]
    if nrows > 1:
        style_reqs.append({"setBasicFilter": {"filter": {
            "range": {"sheetId": new_sheet_id, "startRowIndex": 0,
                      "endRowIndex": nrows,
                      "startColumnIndex": 0, "endColumnIndex": ncols}}}})
    sh.spreadsheets().batchUpdate(
        spreadsheetId=spreadsheet_id, body={"requests": style_reqs}).execute()
    return (f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}"
            f"#gid={new_sheet_id}")


def parse_sheet_id(s):
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", s or "")
    return m.group(1) if m else (s or "").strip()


def append_consolidated(spreadsheet_id, rows, ev, eid, tab_name):
    """Append enriched rows into a user-specified sheet, consolidated schema.

    One shared tab holds every event's rows:
        Event Name | Event URL | Guest Name | LinkedIn Profile |
        Company/Headline | Timestamp
    The tab is created on first use (with a header); later events append
    beneath it. This is the schema the batch GUI targets (event-agnostic,
    timestamped rows) - distinct from the legacy master-tab shape
    (Name, LinkedIn, Location, LA Area?, Confidence, Notes).
    Returns the URL of the tab, with #gid.
    """
    from google.oauth2.credentials import Credentials
    from googleapiclient.discovery import build

    token = os.path.join(HERMES, "google_token.json")
    if not os.path.isfile(token):
        die(f"not authenticated with Google (no {token}). "
            "Run the google-workspace setup first.")
    creds = Credentials.from_authorized_user_file(token)
    sh = build("sheets", "v4", credentials=creds)

    tab = (tab_name or "Event")[:100].strip() or "Event"
    meta = sh.spreadsheets().get(spreadsheetId=spreadsheet_id).execute()
    gid = None
    for s in meta["sheets"]:
        if s["properties"]["title"] == tab:
            gid = s["properties"]["sheetId"]
            break
    if gid is None:
        resp = sh.spreadsheets().batchUpdate(
            spreadsheetId=spreadsheet_id,
            body={"requests": [{"addSheet": {"properties": {"title": tab}}}]}).execute()
        gid = resp["replies"][0]["addSheet"]["properties"]["sheetId"]
        sh.spreadsheets().values().update(
            spreadsheetId=spreadsheet_id, range=f"'{tab}'!A1",
            valueInputOption="RAW",
            body={"values": [["Event Name", "Event URL", "Guest Name",
                              "LinkedIn Profile", "Company/Headline", "Timestamp"]]}).execute()

    stamp = datetime.now().isoformat(timespec="seconds")
    ev_url = f"https://partiful.com/e/{eid}"
    vals = [[ev.get("title") or eid, ev_url,
             r["name"], r.get("url") or "", r.get("headline") or "", stamp]
            for r in rows]
    if vals:
        # Dedupe against rows already in the tab: skip a guest who is already
        # present for THIS event. Keyed on (event URL, normalized guest name)
        # so the same person at a different event still gets added. This makes
        # re-runs idempotent - same as partiful_incremental.py's name matching.
        try:
            existing = sh.spreadsheets().values().get(
                spreadsheetId=spreadsheet_id, range=f"'{tab}'!A:C").execute()
            seen = set()
            for row in existing.get("values", [])[1:]:  # skip header
                if len(row) >= 3 and row[0]:
                    seen.add((row[1].strip(), norm(row[2])))
        except Exception:
            seen = set()  # if reading fails, append everything rather than lose data
        before = len(vals)
        vals = [v for v in vals if (v[1].strip(), norm(v[2])) not in seen]
        if not vals:
            log(f"  {tab!r}: all {before} guests already present - nothing appended")
            return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}#gid={gid}"
        if len(vals) < before:
            log(f"  {tab!r}: {before - len(vals)} of {before} guests already "
                f"present - appending {len(vals)} new")
        sh.spreadsheets().values().append(
            spreadsheetId=spreadsheet_id, range=f"'{tab}'!A:F",
            valueInputOption="RAW", body={"values": vals}).execute()
    return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}#gid={gid}"


def run_pipeline(a):
    """One event, end to end. `a` is an argparse Namespace (see main()).

    Raises PipelineError on failure instead of exiting, so a GUI, cron job or
    test can catch it. Returns a dict with the rows, drops and any sheet URLs.
    Live output goes to stdout AND to LOG (if set); enrichment ticks go to
    PROGRESS(done, total) if set.
    """
    if not os.path.isfile(AUTH):
        die(f"auth file not found: {AUTH}")
    cfg = json.load(open(AUTH, encoding="utf-8"))
    access, uid = refresh_access(cfg)

    m = re.search(r"/e/([A-Za-z0-9]+)", a.event)
    eid = m.group(1) if m else a.event.strip()

    info = api("getEventInfo", {"eventId": eid}, access, uid)
    ev = info.get("event", {})
    guests = api("getGuests", {"eventId": eid}, access, uid)
    log(f"{ev.get('title')}\n  event: {eid}")

    names = [(g.get("user") or {}).get("name") or g.get("name") for g in guests]
    masked = sum(1 for n in names if n and (MASKED.match(n) or re.search(r"\bx{3,}\b", n)))
    if masked:
        die(f"{masked}/{len(names)} names are MASKED - auth is not effective.")

    # Always persist the raw pull: it is the only stable record of guest ids, and
    # a later union across pulls is impossible without it. Callers can override
    # the location, but there is no way to end up with none.
    raw_path = a.raw_out or os.path.join(HERMES, f"partiful-raw-{eid}.json")
    json.dump({"result": {"data": guests}}, open(raw_path, "w", encoding="utf-8"))
    log(f"  raw pull: {raw_path}")
    if a.union_raw and os.path.isfile(a.union_raw):
        prev = json.load(open(a.union_raw, encoding="utf-8"))["result"]["data"]
        seen = {g["id"] for g in guests}
        added = [g for g in prev if g["id"] not in seen]
        guests = guests + added
        log(f"  union: +{len(added)} records from {os.path.basename(a.union_raw)}")

    if a.statuses:
        keep = {s.strip().upper() for s in a.statuses.split(",")}
        present = Counter((g.get("status") or "").upper() for g in guests)
        kept = [g for g in guests if (g.get("status") or "").upper() in keep]
        if not kept:
            # Filtering to a status the event does not use (e.g. --statuses GOING
            # on an approval-gated event where everyone is APPROVED) would
            # otherwise write an empty workbook and look like a quiet success.
            die(f"--statuses {sorted(keep)} matched 0 of {len(guests)} guests. "
                f"Statuses present: {dict(present)}. Drop --statuses or use one of "
                f"those values.")
        if len(kept) < len(guests):
            log(f"  status filter {sorted(keep)}: {len(guests)} -> {len(kept)} "
                  f"(present: {dict(present)})")
        guests = kept

    passing, dropped = classify(guests)
    log(f"  returned: {len(guests)}"
          + ("  <-- exactly 500, likely a CAP" if len(guests) == 500 else ""))
    log(f"  passed:   {len(passing)}   dropped: {len(dropped)}")

    extra = []
    if len(guests) == 500:
        extra.append(("Caveat", "getGuests returned exactly 500 and ignores all pagination "
                                "params; this is a shifting sample, not the full list."))

    if a.no_enrich:
        rows = [{"name": p["name"], "url": None, "loc": None, "flag": "Unknown",
                 "conf": "Low", "note": "", "n": 0} for p in passing]
        write_workbook(rows, dropped, ev, eid, a.out, extra)
        log("wrote:", a.out)
        res = {"rows": rows, "dropped": dropped, "ev": ev, "eid": eid,
               "store": {}, "raw_path": raw_path, "blanked": [], "flagged": [],
               "sheet_url": None, "master_url": None, "target_url": None, "out": a.out}
        if a.upload:
            res["sheet_url"] = upload_to_sheets(a.out, a.sheet_title)
            log("sheet:", res["sheet_url"])
        if a.append_master:
            tab = re.sub(r"[^A-Za-z0-9 _-]+", "", ev.get("title") or eid)[:80].strip() or "Event"
            res["master_url"] = append_to_master(
                rows, tab, parse_sheet_id(a.append_master_sheet) if a.append_master_sheet else MASTER_ID)
            log("master:", res["master_url"])
        if a.append_sheet:
            res["target_url"] = append_consolidated(
                parse_sheet_id(a.append_sheet), rows, ev, eid,
                a.append_tab or (re.sub(r"[^A-Za-z0-9 _-]+", "",
                                        ev.get("title") or eid)[:80].strip() or "Event"))
            log("target sheet:", res["target_url"])
        return res

    key = load_exa_key()
    nm = [p["name"] for p in passing]
    log(f"  enriching {len(nm)} names via Exa ({a.workers} workers)...")
    store, secs = enrich(key, nm, a.workers, a.city, on_progress=PROGRESS)
    log(f"  enrichment done in {secs:.0f}s")
    if a.hits_out:
        json.dump(store, open(a.hits_out, "w", encoding="utf-8"))
        log(f"  hits: {a.hits_out}")

    rows = build_rows(store)
    bad = audit(rows, store)
    if bad:
        die(f"AUDIT FAILED: {len(bad)} URL(s) not backed by a literal /in/ hit: {bad[:5]}")
    log(f"  audit: 0 fabricated URLs across {sum(1 for r in rows if r['url'])} URL cells")
    if a.strict_identity:
        blanked = strict_identity(rows, store)
        if blanked:
            log(f"  strict identity: blanked {len(blanked)} URL(s) with no surname "
                  f"evidence: {blanked}")
    flagged = [r for r in rows if r["url"] and r["note"].startswith("NAME CHECK:")]
    log(f"  identity audit: {len(flagged)}/{sum(1 for r in rows if r['url'])} URL rows "
          f"flagged (profile title lacks the name) - review, do not ship blind")
    for r in flagged:
        log(f"    ~ {r['name']} -> {r['url']}")

    write_workbook(rows, dropped, ev, eid, a.out, extra)
    conf = Counter(r["conf"] for r in rows)
    log(f"  LA=Yes:   {sum(1 for r in rows if r['flag'] == 'Yes')}")
    log(f"  with URL: {sum(1 for r in rows if r['url'])}")
    log(f"  conf:     {conf['High']} High / {conf['Medium']} Medium / {conf['Low']} Low")
    log("wrote:", a.out)
    res = {"rows": rows, "dropped": dropped, "ev": ev, "eid": eid,
           "store": store, "raw_path": raw_path, "blanked": blanked,
           "flagged": flagged, "sheet_url": None, "master_url": None,
           "target_url": None, "out": a.out}
    if a.upload:
        res["sheet_url"] = upload_to_sheets(a.out, a.sheet_title)
        log("sheet:", res["sheet_url"])
    if a.append_master:
        tab = re.sub(r"[^A-Za-z0-9 _-]+", "", ev.get("title") or eid)[:80].strip() or "Event"
        res["master_url"] = append_to_master(
            rows, tab, parse_sheet_id(a.append_master_sheet) if a.append_master_sheet else MASTER_ID)
        log("master:", res["master_url"])
    if a.append_sheet:
        res["target_url"] = append_consolidated(
            parse_sheet_id(a.append_sheet), rows, ev, eid,
            a.append_tab or (re.sub(r"[^A-Za-z0-9 _-]+", "",
                                    ev.get("title") or eid)[:80].strip() or "Event"))
        log("target sheet:", res["target_url"])
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--event", required=True, help="Partiful event URL or ID")
    ap.add_argument("--out", required=True, help="output .xlsx path")
    ap.add_argument("--city", default="Los Angeles", help="city used in the search query")
    ap.add_argument("--workers", type=int, default=8, help="parallel Exa searches")
    ap.add_argument("--statuses", default="", help="comma list to keep, e.g. GOING,MAYBE")
    ap.add_argument("--raw-out", default=None, help="save the raw guest pull here")
    ap.add_argument("--upload", action="store_true",
                    help="also create a Google Sheet from the workbook and print its URL")
    ap.add_argument("--sheet-title", default=None,
                    help="title for the created Google Sheet (default: the workbook name)")
    ap.add_argument("--union-raw", default=None,
                    help="merge a previous raw pull (recovers past the 500 cap)")
    ap.add_argument("--no-enrich", action="store_true", help="names only, skip LinkedIn")
    ap.add_argument("--append-master", action="store_true",
                    help="add a tab to the Oxen-Forecast master Google Sheet")
    ap.add_argument("--append-master-sheet", default=None,
                    help="with --append-master: target a different sheet (URL or ID) "
                         "instead of the Oxen-Forecast master")
    ap.add_argument("--append-sheet", default=None,
                    help="append rows into a user-specified Google Sheet (URL or ID); "
                         "creates/uses ONE consolidated tab (see --append-tab)")
    ap.add_argument("--append-tab", default=None,
                    help="tab name for --append-sheet (default: the event title)")
    ap.add_argument("--strict-identity", action="store_true",
                    help="blank URLs whose matched profile shows no surname evidence")
    ap.add_argument("--hits-out", default=None,
                    help="save the raw Exa hit store here so the audits can be re-run offline")
    a = ap.parse_args()
    run_pipeline(a)


if __name__ == "__main__":
    main()
